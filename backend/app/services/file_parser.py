"""Lightweight document-to-text parsing. No OCR — matches the assignment's
explicit "production-grade OCR not required" scope. Scanned/image-only PDFs
will yield little/no text, which is an accepted limitation for this project.
"""

import io
from email import message_from_bytes, policy

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import DocumentParseError

SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "message/rfc822": "eml",
}


def guess_kind(filename: str, mime_type: str | None) -> str:
    lower = (filename or "").lower()
    if lower.endswith(".pdf"):
        return "pdf"
    if lower.endswith(".docx"):
        return "docx"
    if lower.endswith(".eml"):
        return "eml"
    if lower.endswith(".txt"):
        return "txt"
    if mime_type in SUPPORTED_MIME_TYPES:
        return SUPPORTED_MIME_TYPES[mime_type]
    raise DocumentParseError(f"Unsupported file type for '{filename}' ({mime_type})")


def parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip()
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Failed to parse PDF: {exc}") from exc


def parse_docx(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs).strip()
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Failed to parse DOCX: {exc}") from exc


def parse_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def parse_eml(data: bytes) -> str:
    try:
        msg = message_from_bytes(data, policy=policy.default)
        header_lines = []
        for header in ("From", "To", "Subject", "Date"):
            if msg[header]:
                header_lines.append(f"{header}: {msg[header]}")

        body_parts = []
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain" and not part.get_filename():
                    body_parts.append(part.get_content())
        else:
            body_parts.append(msg.get_content())

        return "\n".join(header_lines) + "\n\n" + "\n".join(body_parts)
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Failed to parse EML: {exc}") from exc


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "txt": parse_txt,
    "eml": parse_eml,
}


def extract_text(filename: str, mime_type: str | None, data: bytes) -> str:
    kind = guess_kind(filename, mime_type)
    text = PARSERS[kind](data)
    if not text.strip():
        raise DocumentParseError(
            f"No extractable text found in '{filename}'. Scanned/image-only documents are not supported (no OCR)."
        )
    return text
